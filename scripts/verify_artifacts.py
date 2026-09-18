"""Live model + browser export test. Requires playwright, pypdf, openpyxl.

python3 scripts/verify_artifacts.py --url http://127.0.0.1:5175/ai
"""
import argparse
import hashlib
import json
from pathlib import Path
import tempfile
from playwright.sync_api import sync_playwright, expect
from pypdf import PdfReader
from openpyxl import load_workbook

parser = argparse.ArgumentParser()
parser.add_argument('--url', default='http://127.0.0.1:5175/ai')
parser.add_argument('--model', default='deepseek-chat')
args = parser.parse_args()
output_dir = Path(tempfile.mkdtemp(prefix='agent-export-qa-'))

with sync_playwright() as p:
    browser = p.chromium.launch(headless=True)
    try:
        page = browser.new_page(accept_downloads=True)
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.goto(args.url)
        page.wait_for_load_state('networkidle')
        page.evaluate('(model) => localStorage.setItem("selectedModelId", model)', args.model)
        page.reload()
        page.wait_for_load_state('networkidle')
        page.get_by_role('button', name='智能体模式', exact=True).click()
        page.locator('textarea').fill('请实际生成两个可下载文件：1.PDF标题“项目预算说明”，正文写中文说明：设计费用1200元，开发费用3400元；2.Excel文件标题“项目预算”，工作表“预算”，列名“项目”和“金额（元）”，两行数据“设计”1200和“开发”3400，金额列添加SUM公式合计。最后简短回答。')
        page.locator('textarea').press('Enter')
        expect(page.locator('.agent-artifact')).to_have_count(2, timeout=180000)
        expect(page.locator('.agent-trace--completed')).to_be_visible(timeout=180000)
        artifacts = page.evaluate('JSON.parse(localStorage.getItem("messages"))[0].data.at(-1).artifacts')
        assert {file['format'] for file in artifacts} == {'pdf', 'xlsx'}
        digests = {}
        for file in artifacts:
            with page.expect_download() as download_info:
                page.get_by_role('button', name=f'下载 {file["fileName"]}', exact=True).click()
            download = download_info.value
            assert download.failure() is None
            target = output_dir / download.suggested_filename
            download.save_as(target)
            digests[file['fileName']] = hashlib.sha256(target.read_bytes()).hexdigest()
            if file['format'] == 'pdf':
                text = '\n'.join(pdf_page.extract_text() for pdf_page in PdfReader(target).pages)
                assert '设计' in text and '开发' in text and '1200' in text and '3400' in text, text
            else:
                formulas = load_workbook(target, data_only=False)
                values = load_workbook(target, data_only=True)
                sheet = formulas['预算']
                assert sheet['B2'].value == 1200 and sheet['B3'].value == 3400
                assert sheet['B4'].value == '=SUM(B2:B3)'
                assert values['预算']['B4'].value == 4600
                formulas.close()
                values.close()
            print(f'PASS: downloaded and parsed {file["fileName"]}', flush=True)
        page.reload()
        page.wait_for_load_state('networkidle')
        expect(page.locator('.agent-artifact')).to_have_count(2)
        page.locator('.agent-artifacts').scroll_into_view_if_needed()
        page.screenshot(path=str(output_dir / 'cards.png'), full_page=True)
        for file in artifacts:
            with page.expect_download() as download_info:
                page.get_by_role('button', name=f'下载 {file["fileName"]}', exact=True).click()
            download = download_info.value
            assert download.failure() is None
            assert hashlib.sha256(Path(download.path()).read_bytes()).hexdigest() == digests[file['fileName']]
        page.locator('textarea').fill('请给出刚才两个文件的Markdown下载链接，沿用原来的地址，不重新生成文件。')
        page.locator('textarea').press('Enter')
        expect(page.locator('.agent-trace--completed')).to_have_count(2, timeout=150000)
        links = page.locator('.message.bot').last.locator('a[href*="/api/files/"]')
        expect(links).to_have_count(2)
        for file in artifacts:
            link = page.locator('.message.bot').last.locator(f'a[href*="{file["fileId"]}"]')
            expect(link).to_have_count(1)
            with page.expect_download() as download_info:
                link.click()
            download = download_info.value
            assert download.failure() is None
            assert hashlib.sha256(Path(download.path()).read_bytes()).hexdigest() == digests[file['fileName']]
        assert not errors, errors
        page.screenshot(path=str(output_dir / 'downloads.png'), full_page=True)
        (output_dir / 'artifacts.json').write_text(json.dumps(artifacts, ensure_ascii=False, indent=2))
        print(f'PASS: downloads survive refresh; no browser errors. Evidence: {output_dir}', flush=True)
    finally:
        browser.close()
